# backend/services/parser.py
"""
TASK-002: ドキュメント解析モジュール (タイムアウトガードレール付き)

対応フォーマット:
  - PDF  (.pdf)  : PyMuPDF で各ページのテキストを抽出
  - Word (.docx) : python-docx で段落テキストを抽出
  - テキスト (.txt / .md) : UTF-8 デコードそのまま返却

制約:
  - 最大ファイルサイズ: 20 MB（MAX_UPLOAD_SIZE）
  - LLM保護用の最大抽出文字数: 5,000文字（MAX_EXTRACT_CHARS）
"""

import io
import os
from pathlib import Path
from typing import Optional, Tuple

import fitz  # PyMuPDF
from docx import Document
from pydantic import BaseModel

MAX_UPLOAD_SIZE: int = int(os.getenv("MAX_UPLOAD_SIZE", 20 * 1024 * 1024))  # 20 MB

# 💡 【タイムアウト対策】LLM（Ollama）のパンクを防ぐ最大文字数の制限値
MAX_EXTRACT_CHARS: int = int(os.getenv("MAX_EXTRACT_CHARS", 5000))


class UnsupportedFileTypeError(ValueError):
    """サポートされていないファイル形式"""


class FileTooLargeError(ValueError):
    """ファイルサイズ超過"""


class ParsedDocument(BaseModel):
    """`POST /api/documents/parse`のレスポンス（Backend依頼_仕様書ファイル変換API.md）"""

    filename: str
    text: str
    char_count: int
    page_count: Optional[int] = None
    truncated: bool = False


def parse_document(filename: str, file_bytes: bytes) -> str:
    _validate_size(file_bytes)
    ext = Path(filename).suffix.lower()

    # 各フォーマットからテキストを抽出
    if ext == ".pdf":
        text, _ = _extract_pdf(file_bytes)
    elif ext in (".docx",):
        text = _extract_docx(file_bytes)
    elif ext in (".txt", ".md"):
        text = _extract_text(file_bytes)
    else:
        raise UnsupportedFileTypeError(
            f"未対応のファイル形式です: {ext}。"
            "対応形式: PDF / Word (.docx) / テキスト (.txt, .md)"
        )

    # 💡 【バグ・タイムアウト対策】抽出テキストが空、または画像PDFだった場合のガード
    if not text.strip():
        raise ValueError(
            f"ファイル「{filename}」からテキストを抽出できませんでした。"
            "中身が空か、スキャンされた画像ベースのファイルである可能性があります。"
        )

    # 💡 【重要】Ollamaの処理負荷を下げ、504エラーを防ぐための文字数カット
    if len(text) > MAX_EXTRACT_CHARS:
        text = text[:MAX_EXTRACT_CHARS] + "\n\n... [システム警告: 仕様書が長すぎるため、これ以降のテキストはタイムアウト防止のため省略されました] ..."

    return text


def parse_document_full(filename: str, file_bytes: bytes) -> ParsedDocument:
    """`POST /api/documents/parse`用（Backend依頼_仕様書ファイル変換API.md 案A）。

    `parse_document()`と異なり、`MAX_EXTRACT_CHARS`による切り詰めを行わない
    ——Requirementsステージは文書をチャンク単位で処理するため、ここで事前に
    切り詰める必要が無い（案Aの理由）。上限は`_validate_size()`による
    `MAX_UPLOAD_SIZE`（ファイルサイズ）のみで、`truncated`は常に`False`。
    PDFのページ数も返す（Word/テキストは`None`）。
    """
    _validate_size(file_bytes)
    ext = Path(filename).suffix.lower()

    page_count: Optional[int] = None
    if ext == ".pdf":
        text, page_count = _extract_pdf(file_bytes)
    elif ext == ".docx":
        text = _extract_docx(file_bytes)
    elif ext in (".txt", ".md"):
        text = _extract_text(file_bytes)
    else:
        raise UnsupportedFileTypeError(
            f"未対応のファイル形式です: {ext}。"
            "対応形式: PDF / Word (.docx) / テキスト (.txt, .md)"
        )

    if not text.strip():
        raise ValueError(
            f"ファイル「{filename}」からテキストを抽出できませんでした。"
            "中身が空か、スキャンされた画像ベースのファイルである可能性があります。"
        )

    return ParsedDocument(
        filename=filename,
        text=text,
        char_count=len(text),
        page_count=page_count,
        truncated=False,
    )


def _validate_size(file_bytes: bytes) -> None:
    size = len(file_bytes)
    if size > MAX_UPLOAD_SIZE:
        mb = MAX_UPLOAD_SIZE // (1024 * 1024)
        raise FileTooLargeError(
            f"ファイルサイズ ({size / 1024 / 1024:.1f} MB) が上限 {mb} MB を超えています。"
        )


def _extract_pdf(file_bytes: bytes) -> Tuple[str, int]:
    doc = fitz.open(stream=file_bytes, filetype="pdf")
    pages = [page.get_text() for page in doc]
    page_count = len(doc)
    doc.close()
    return "\n".join(pages).strip(), page_count


def _extract_docx(file_bytes: bytes) -> str:
    doc = Document(io.BytesIO(file_bytes))
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    return "\n".join(paragraphs).strip()


def _extract_text(file_bytes: bytes) -> str:
    return file_bytes.decode("utf-8", errors="replace").strip()